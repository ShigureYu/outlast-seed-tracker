"""把 data/seed.xlsx 里的各关答案表导入 sqlite。

每一关用一条 TrialSpec 描述：读哪个 sheet、表名（Trial ID）、中文名、
列定义、以及“码长”和“是否需要拆分”。

两种表格形态：

1. 直接多列（被告判决）：每个图案一列，每格一个 3 位码。
2. 单列合并（度假村）：xlsx 里存一个 3 位数（如 743），
   需要拆成 3 个 1 位码分别存入 killer / weapon / room。

约定：
- 表名一律用 Trial ID（大写），列名 ASCII。
- 码值一律 TEXT，保住 '014' 这种前导零。
- 源表缺格/脏格存 NULL，不整行丢弃。
"""

from __future__ import annotations

import json
import logging
import re
import sqlite3
from contextlib import closing
from dataclasses import dataclass
from pathlib import Path

import openpyxl

logger = logging.getLogger(__name__)

from app_paths import DATA_DIR, XLSX_PATH, enable_utf8_stdout

DB_PATH = DATA_DIR / "seed.db"  # 正式库，固定名字，故意不跟随 config.json
META_TABLE = "TRIAL_META"

TRIAL_ID_RE = re.compile(r"^[A-Z][A-Z0-9]*(?:_[A-Z0-9]+)*$")
# 表头里的图案名带 emoji（骷髅☠️），入库的显示名只要中文
NON_CJK_RE = re.compile(r"[^\u4e00-\u9fff]")


@dataclass(frozen=True)
class TrialSpec:
    sheet: str | None  # None = 只建表登记结构，不导入数据
    trial_id: str
    name_zh: str
    map_name: str
    columns: tuple[tuple[str, str], ...]  # (列名, 显示名)
    code_length: int = 3
    split_from: int | None = None  # 非空时表示：这一列是合并码，需要拆开

    @property
    def column_names(self) -> tuple[str, ...]:
        return tuple(name for name, _ in self.columns)


TRIALS: tuple[TrialSpec, ...] = (
    TrialSpec(
        sheet="被告判决",
        trial_id="CHJ_MT02",
        name_zh="对被告作出判决",
        map_name="CourthouseJudicial",
        columns=(
            ("skull", "骷髅"),
            ("spiral", "螺旋"),
            ("triangle", "三角形"),
            ("star", "五角星"),
            ("sun", "太阳"),
        ),
        code_length=3,
    ),
    TrialSpec(
        sheet="度假村",
        trial_id="RES_MT01",
        name_zh="侦破谋杀案",
        map_name="ResortSpa",
        columns=(
            ("killer", "凶手"),
            ("weapon", "凶器"),
            ("room", "房间"),
        ),
        code_length=1,
        split_from=1,  # xlsx 第 2 列是一个三位数，拆成三段
    ),
    TrialSpec(
        sheet=None,  # 暂无数据表：只登记结构，让界面能录入
        trial_id="MHS_MT01",
        name_zh="清空保险库",
        map_name="MobsterHideoutSyndicate",
        columns=(
            ("triangle", "三角形"),
            ("skull", "骷髅"),
            ("sun", "太阳"),
        ),
        code_length=1,
    ),
)


def _clean_label(text: str) -> str:
    """'骷髅☠️' -> '骷髅'：去掉 emoji 和符号，只留中文。"""
    return NON_CJK_RE.sub("", text) or text


def _normalize_seed(raw: object) -> int | None:
    if raw is None:
        return None
    text = str(raw).strip()
    if not text.isdigit():
        return None
    seed = int(text)
    return seed if 1 <= seed <= 999 else None


def _parse_codes(spec: TrialSpec, row: tuple) -> list[str | None]:
    """从一行里取出该关所有码值；不合法/缺失的位置返回 None。"""
    if spec.split_from is None:
        width = len(spec.columns)
        raw_values = row[1 : 1 + width]
    else:
        joined = row[spec.split_from]
        text = str(joined).strip() if joined is not None else ""
        width = len(spec.columns)
        # 合并码必须正好是 width 位数字，否则整行不拆
        raw_values = list(text) if re.fullmatch(rf"\d{{{width}}}", text) else [None] * width

    result: list[str | None] = []
    for value in raw_values:
        text = str(value).strip() if value is not None else ""
        result.append(text if re.fullmatch(rf"\d{{{spec.code_length}}}", text) else None)
    return result


def create_schema(conn: sqlite3.Connection) -> None:
    conn.execute(
        f"""
        CREATE TABLE IF NOT EXISTS {META_TABLE} (
            trial_id    TEXT PRIMARY KEY,
            name_zh     TEXT NOT NULL,
            map_name    TEXT,
            labels      TEXT,
            code_length INTEGER DEFAULT 3
        )
        """
    )
    meta_columns = [row[1] for row in conn.execute(f"PRAGMA table_info({META_TABLE})")]
    if "labels" not in meta_columns:
        conn.execute(f"ALTER TABLE {META_TABLE} ADD COLUMN labels TEXT")
    if "code_length" not in meta_columns:
        conn.execute(f"ALTER TABLE {META_TABLE} ADD COLUMN code_length INTEGER DEFAULT 3")


def _table_columns(conn: sqlite3.Connection, table: str) -> list[str]:
    return [row[1] for row in conn.execute(f"PRAGMA table_info({table})")]


def ensure_table(conn: sqlite3.Connection, spec: TrialSpec) -> None:
    if not TRIAL_ID_RE.match(spec.trial_id):
        raise ValueError(f"Trial ID 不合法: {spec.trial_id}")

    definition = ",\n            ".join(f"{name} TEXT" for name in spec.column_names)
    conn.execute(
        f"""
        CREATE TABLE IF NOT EXISTS {spec.trial_id} (
            seed INTEGER PRIMARY KEY,
            {definition},
            updated_at TEXT
        )
        """
    )
    if "updated_at" not in _table_columns(conn, spec.trial_id):
        conn.execute(f"ALTER TABLE {spec.trial_id} ADD COLUMN updated_at TEXT")


def load_rows(spec: TrialSpec) -> tuple[list[tuple], list[tuple[int, int, str]]]:
    workbook = openpyxl.load_workbook(XLSX_PATH, data_only=True, read_only=True)
    if spec.sheet not in workbook.sheetnames:
        workbook.close()
        raise SystemExit(f"找不到 sheet: {spec.sheet}（现有: {workbook.sheetnames}）")

    sheet = workbook[spec.sheet]
    rows: list[tuple] = []
    gaps: list[tuple[int, int, str]] = []
    width = len(spec.columns)

    for index, row in enumerate(sheet.iter_rows(values_only=True), start=1):
        if index == 1:
            logger.info("[%s] 表头: %s", spec.sheet, [c for c in row if c is not None])
            continue
        if row is None or row[0] is None:
            continue

        seed = _normalize_seed(row[0])
        if seed is None:
            logger.warning("[%s] 第 %d 行 seed 非法，跳过: %r", spec.sheet, index, row[0])
            continue

        codes = _parse_codes(spec, row)
        missing = [
            f"{label}={value!r}"
            for (name, label), value in zip(spec.columns, codes)
            if value is None
        ]
        if missing:
            gaps.append((index, seed, ", ".join(missing)))

        rows.append((seed, *codes))

    workbook.close()
    return rows, gaps


def register_only(conn: sqlite3.Connection, spec: TrialSpec) -> None:
    """没有数据表时只建空表 + 写 meta，界面就能按标签渲染输入框。"""
    ensure_table(conn, spec)
    labels = {name: label for name, label in spec.columns}
    conn.execute(
        f"INSERT OR REPLACE INTO {META_TABLE} (trial_id, name_zh, map_name, labels, code_length) "
        f"VALUES (?, ?, ?, ?, ?)",
        (
            spec.trial_id,
            spec.name_zh,
            spec.map_name,
            json.dumps(labels, ensure_ascii=False),
            spec.code_length,
        ),
    )
    conn.commit()
    total = conn.execute(f"SELECT COUNT(*) FROM {spec.trial_id}").fetchone()[0]
    print(f"[{spec.trial_id}] {spec.name_zh}: 仅登记结构，表内现有 {total} 行")


def import_trial(conn: sqlite3.Connection, spec: TrialSpec) -> None:
    if spec.sheet is None:
        register_only(conn, spec)
        return

    rows, gaps = load_rows(spec)

    seen: dict[int, int] = {}
    for seed, *_ in rows:
        seen[seed] = seen.get(seed, 0) + 1
    duplicated = [s for s, n in seen.items() if n > 1]
    if duplicated:
        logger.warning("[%s] 有 %d 个 seed 重复: %s", spec.trial_id, len(duplicated), duplicated[:10])

    ensure_table(conn, spec)
    labels = {name: label for name, label in spec.columns}
    conn.execute(
        f"INSERT OR REPLACE INTO {META_TABLE} (trial_id, name_zh, map_name, labels, code_length) "
        f"VALUES (?, ?, ?, ?, ?)",
        (
            spec.trial_id,
            spec.name_zh,
            spec.map_name,
            json.dumps(labels, ensure_ascii=False),
            spec.code_length,
        ),
    )
    conn.executemany(
        f"INSERT OR REPLACE INTO {spec.trial_id} (seed, {', '.join(spec.column_names)}) "
        f"VALUES ({', '.join('?' * (len(spec.columns) + 1))})",
        rows,
    )
    conn.commit()

    total = conn.execute(f"SELECT COUNT(*) FROM {spec.trial_id}").fetchone()[0]
    incomplete = conn.execute(
        f"SELECT COUNT(*) FROM {spec.trial_id} WHERE "
        + " OR ".join(f"{name} IS NULL" for name in spec.column_names)
    ).fetchone()[0]

    print(
        f"[{spec.trial_id}] {spec.name_zh}: 入库 {len(rows)} 行，"
        f"其中 {incomplete} 行有空格子，表内共 {total} 行"
    )
    for index, seed, reason in gaps[:10]:
        print(f"  第 {index} 行 seed={seed} 存为 NULL: {reason}")


def import_all() -> None:
    if not XLSX_PATH.is_file():
        raise SystemExit(f"找不到表格: {XLSX_PATH}")

    with closing(sqlite3.connect(DB_PATH)) as conn:
        create_schema(conn)
        for spec in TRIALS:
            import_trial(conn, spec)

    print(f"数据库: {DB_PATH}")


def main() -> None:
    enable_utf8_stdout()
    logging.basicConfig(level=logging.INFO, format="%(levelname)s: %(message)s")
    import_all()


if __name__ == "__main__":
    main()
